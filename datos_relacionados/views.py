from rest_framework.viewsets import ModelViewSet
from .serializers import CursoActivoSerializer, HorarioSerializer
from .models import CursoActivo, Horario
from datos_simples.models import DistribucionCargaAcademica, Curso, Carrera, Docente

from django.http import JsonResponse
from rest_framework.decorators import api_view
from openpyxl import load_workbook
from datos_simples.views import retrieveAll, agregarFilas, atributoADataframe
import pandas as pd
from sqlalchemy import create_engine
import datetime

url = 'postgresql+psycopg2://roswell:12345@localhost:5432/ingenieria_software'
engine = create_engine(url=url, echo=False)

recuperar_id = lambda objeto : objeto.id
recuperar_nombre = lambda objeto : objeto.nombre
recuperar_codigo = lambda objeto : objeto.codigo
to_time = lambda entero:datetime.time(entero)

# leer distribucion de carga academica 2022
def leerDistribucion():
  archivos = retrieveAll(DistribucionCargaAcademica)
  archivo = archivos[0]
  archivo = archivo.archivo
  # leer el workbook
  data_excel = load_workbook(archivo, data_only=True)
  hoja_distribucion = data_excel['DISTR. SEM 2022-2-INF']
  data = hoja_distribucion.tables['CARGA_ACAD']
  data = hoja_distribucion[data.ref]
  lista_filas = []
  for fila in data:
    atributos = []
    for atributo in fila:
      atributos.append(atributo.value)
    lista_filas.append(atributos)
  dataframe = pd.DataFrame(data=lista_filas[1:], index=None, columns=lista_filas[0])
  dataframe['CARRERA'] = agregarFilas(dataframe['CARRERA'])
  dataframe['CODIGO'] = agregarFilas(dataframe['CODIGO'])
  dataframe = dataframe.explode(['CODIGO', 'CARRERA'])
  return dataframe

@api_view(['GET'])
def subirCursosActivos(request):
  distribucion = leerDistribucion()
  # agregar cursos activos
  # -- recuperar 3 atributos
  cursos_activos = distribucion[['CODIGO', 'CURSO', 'CARRERA']]
  cursos_activos.drop_duplicates(inplace=True)
  # -- recuperar los cursos de la base de datos
  cursos_db = retrieveAll(Curso)
  cursos_db = atributoADataframe(cursos_db, 'objeto_curso')
  # -- mapear el id
  cursos_db['datos_curso'] = pd.Series(map(recuperar_id, cursos_db['objeto_curso']))
  cursos_db['CURSO'] = pd.Series(map(recuperar_nombre, cursos_db['objeto_curso']))
  cursos_db = cursos_db[['datos_curso', 'CURSO']]
  # -- recuperar las carreras de la base de datos
  carreras_db = retrieveAll(Carrera)
  carreras_db = atributoADataframe(carreras_db, 'objeto_carrera')
  carreras_db['datos_carrera'] = pd.Series(map(recuperar_id, carreras_db['objeto_carrera']))
  carreras_db['CARRERA'] = pd.Series(map(recuperar_nombre, carreras_db['objeto_carrera']))
  carreras_db = carreras_db[['datos_carrera', 'CARRERA']]
  # -- unir con los cursos activos
  cursos_activos = cursos_activos.join(cursos_db.set_index('CURSO'), on='CURSO')
  cursos_activos = cursos_activos.join(carreras_db.set_index('CARRERA'), on='CARRERA')
  # -- quedarme con 3 atributos
  cursos_activos = cursos_activos[['CODIGO', 'datos_curso', 'datos_carrera']]
  cursos_activos.columns = ['codigo', 'datos_curso_id', 'datos_carrera_id']
  with engine.connect().execution_options(autocommit=True) as conn:
      cursos_activos.to_sql(CursoActivo._meta.db_table, con=conn, index=False, if_exists='append', method='multi')
  return JsonResponse({'msg': 'correcto'})

@api_view(['GET'])
def subirHorarios(request):
  distribucion = leerDistribucion()
  horarios = distribucion[['CODIGO', 'CRED.', 'TIPO', 'GPO',
    'HT', 'HP', 'HR/\nINICIO', 'HR/\nFIN', 'DIA', 'AULA', 'DOCENTES']]
  index_remove = horarios[ (horarios['DOCENTES'] == 'GRUPO DESACTIVADO') | (horarios['DOCENTES'] == 'CURSO O GRUPO POR DESACTIVAR') | (horarios['DOCENTES'] == 'CURSO DESACTIVADO') ].index
  horarios.drop(index_remove, inplace=True)
  # recuperar cursos activos
  cursos_activos_db = retrieveAll(CursoActivo)
  cursos_activos_db = atributoADataframe(cursos_activos_db, 'objeto_c_activo')
  cursos_activos_db['curso_activo'] = pd.Series(map(recuperar_id, cursos_activos_db['objeto_c_activo']))
  cursos_activos_db['CODIGO'] = pd.Series(map(recuperar_codigo, cursos_activos_db['objeto_c_activo']))
  # recuperar docentes
  docentes_db = retrieveAll(Docente)
  docentes_db = atributoADataframe(docentes_db, 'objeto_docente')
  docentes_db['docente'] = pd.Series(map(recuperar_id, docentes_db['objeto_docente']))
  docentes_db['DOCENTES'] = pd.Series(map(recuperar_nombre, docentes_db['objeto_docente']))
  # unir con los horarios
  horarios = horarios.join(cursos_activos_db.set_index('CODIGO'), on='CODIGO')
  horarios = horarios.join(docentes_db.set_index('DOCENTES'), on='DOCENTES')
  horarios = horarios[['curso_activo', 'CRED.', 'TIPO', 'GPO',
    'HT', 'HP', 'HR/\nINICIO', 'HR/\nFIN', 'DIA', 'AULA', 'docente']]
  horarios.columns = ['curso_activo_id', 'creditos', 'tipo', 'grupo',
    'ht', 'hp', 'hi', 'hf', 'dia', 'aula', 'docente_id']
  horarios['ht'] = pd.Series(map(to_time, horarios['ht']))
  horarios['hp'] = pd.Series(map(to_time, horarios['hp']))
  horarios['hi'] = pd.Series(map(to_time, horarios['hi']))
  horarios['hf'] = pd.Series(map(to_time, horarios['hf']))
  with engine.connect().execution_options(autocommit=True) as conn:
      horarios.to_sql(Horario._meta.db_table, con=conn, index=False, if_exists='append', method='multi')
  return JsonResponse({'msg': 'correcto'})

# Create your views here.
class CursoActivoViewSet(ModelViewSet):
  queryset = CursoActivo.objects.all()
  serializer_class = CursoActivoSerializer

class HorarioViewSet(ModelViewSet):
  queryset = Horario.objects.all()
  serializer_class = HorarioSerializer